"""Export transactions to multiple formats."""

import json
from abc import ABC, abstractmethod
from html import escape
from pathlib import Path
from typing import Optional

from .models import Transacao, Extrato
from .utils import configurar_logging

logger = configurar_logging(__name__)


class ExportadorBase(ABC):
    """Abstract base class for exporters."""
    
    def __init__(self, caminho_saida: Path):
        self.caminho_saida = Path(caminho_saida)
        self.logger = logger
    
    @abstractmethod
    def exportar(self, extrato: Extrato) -> None:
        """Export statement to file.
        
        Args:
            extrato: Extrato object with transactions
        """
        pass


class ExportadorOFX(ExportadorBase):
    """Export transactions to OFX format."""
    
    def exportar(self, extrato: Extrato) -> None:
        """Generate OFX file.
        
        Args:
            extrato: Extrato object
        """
        self.logger.info(f"Gerando OFX: {self.caminho_saida}")
        
        ofx_conteudo = self._gerar_header_ofx()
        ofx_conteudo += self._gerar_transacoes_ofx(extrato.transacoes)
        ofx_conteudo += self._gerar_footer_ofx(extrato)
        
        self.caminho_saida.parent.mkdir(parents=True, exist_ok=True)
        with open(self.caminho_saida, 'w', encoding='utf-8') as arquivo:
            arquivo.write(ofx_conteudo)
        
        self.logger.info(f"Arquivo OFX criado: {self.caminho_saida}")
    
    def _gerar_header_ofx(self) -> str:
        """Generate OFX header."""
        return """OFXHEADER:100
DATA:OFXSGML
VERSION:102
SECURITY:NONE
ENCODING:USASCII
CHARSET:1252
COMPRESSION:NONE
OLDFILEUID:NONE
NEWFILEUID:NONE

<OFX>
  <BANKMSGSRSV1>
    <STMTTRNRS>
      <TRNUID>1</TRNUID>
      <STATUS>
        <CODE>0</CODE>
        <SEVERITY>INFO</SEVERITY>
      </STATUS>
      <STMTRS>
        <CURDEF>BRL</CURDEF>
        <BANKACCTFROM>
          <BANKID>{agencia}</BANKID>
          <ACCTID>{conta}</ACCTID>
          <ACCTTYPE>CHECKING</ACCTTYPE>
        </BANKACCTFROM>
        <BANKTRANLIST>
"""
    
    def _gerar_transacoes_ofx(self, transacoes: list[Transacao]) -> str:
        """Generate OFX transaction lines."""
        ofx = ""
        for i, transacao in enumerate(transacoes):
            tipo = transacao.tipo.value
            ofx += f"""          <STMTTRN>
            <TRNTYPE>{tipo}</TRNTYPE>
            <DTPOSTED>{transacao.data}</DTPOSTED>
            <TRNAMT>{transacao.valor}</TRNAMT>
            <FITID>{transacao.data}_{i}</FITID>
            <MEMO>{escape(transacao.descricao)}</MEMO>
          </STMTTRN>
"""
        return ofx
    
    def _gerar_footer_ofx(self, extrato: Extrato) -> str:
        """Generate OFX footer."""
        saldo = extrato.saldo_final or extrato.saldo_liquido()
        data_fim = extrato.data_fim or "20260830"
        
        return f"""        </BANKTRANLIST>
        <LEDGERBAL>
          <BALAMT>{saldo}</BALAMT>
          <DTASOF>{data_fim}</DTASOF>
        </LEDGERBAL>
      </STMTRS>
    </STMTTRNRS>
  </BANKMSGSRSV1>
</OFX>"""


class ExportadorCSV(ExportadorBase):
    """Export transactions to CSV format."""
    
    def exportar(self, extrato: Extrato) -> None:
        """Generate CSV file.
        
        Args:
            extrato: Extrato object
        """
        self.logger.info(f"Gerando CSV: {self.caminho_saida}")
        
        self.caminho_saida.parent.mkdir(parents=True, exist_ok=True)
        
        with open(self.caminho_saida, 'w', encoding='utf-8', newline='') as arquivo:
            # Header
            arquivo.write("Data,Descrição,Valor,Tipo\n")
            
            # Transactions
            for transacao in extrato.transacoes:
                data_formatada = f"{transacao.data[0:4]}-{transacao.data[4:6]}-{transacao.data[6:8]}"
                descricao_escape = transacao.descricao.replace('"', '""')
                arquivo.write(
                    f"{data_formatada},"
                    f'"{descricao_escape}",'
                    f"{transacao.valor},"
                    f"{transacao.tipo.value}\n"
                )
        
        self.logger.info(f"Arquivo CSV criado: {self.caminho_saida}")


class ExportadorJSON(ExportadorBase):
    """Export transactions to JSON format."""
    
    def exportar(self, extrato: Extrato) -> None:
        """Generate JSON file.
        
        Args:
            extrato: Extrato object
        """
        self.logger.info(f"Gerando JSON: {self.caminho_saida}")
        
        dados = {
            "banco": extrato.banco.value,
            "agencia": extrato.agencia,
            "conta": extrato.conta,
            "data_inicio": extrato.data_inicio,
            "data_fim": extrato.data_fim,
            "total_transacoes": len(extrato.transacoes),
            "total_creditos": str(extrato.total_creditos()),
            "total_debitos": str(extrato.total_debitos()),
            "saldo_liquido": str(extrato.saldo_liquido()),
            "transacoes": [
                {
                    "data": t.data,
                    "descricao": t.descricao,
                    "valor": str(t.valor),
                    "tipo": t.tipo.value,
                }
                for t in extrato.transacoes
            ]
        }
        
        self.caminho_saida.parent.mkdir(parents=True, exist_ok=True)
        with open(self.caminho_saida, 'w', encoding='utf-8') as arquivo:
            json.dump(dados, arquivo, ensure_ascii=False, indent=2)
        
        self.logger.info(f"Arquivo JSON criado: {self.caminho_saida}")


class ExportadorXLSX(ExportadorBase):
    """Export transactions to XLSX format."""
    
    def exportar(self, extrato: Extrato) -> None:
        """Generate XLSX file.
        
        Args:
            extrato: Extrato object
            
        Requires: openpyxl package
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            self.logger.error("openpyxl não está instalado. Use: pip install openpyxl")
            raise
        
        self.logger.info(f"Gerando XLSX: {self.caminho_saida}")
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Extrato"
        
        # Header info
        ws['A1'] = f"Banco: {extrato.banco.value}"
        ws['A2'] = f"Agência: {extrato.agencia}"
        ws['A3'] = f"Conta: {extrato.conta}"
        ws['A4'] = f"Período: {extrato.data_inicio} a {extrato.data_fim}"
        
        # Column headers
        headers = ["Data", "Descrição", "Valor", "Tipo"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        
        # Data
        for row, transacao in enumerate(extrato.transacoes, 7):
            data_formatada = f"{transacao.data[0:4]}-{transacao.data[4:6]}-{transacao.data[6:8]}"
            ws.cell(row=row, column=1).value = data_formatada
            ws.cell(row=row, column=2).value = transacao.descricao
            ws.cell(row=row, column=3).value = float(transacao.valor)
            ws.cell(row=row, column=4).value = transacao.tipo.value
            
            # Format currency column
            ws.cell(row=row, column=3).number_format = '#,##0.00'
        
        # Summary row
        total_row = len(extrato.transacoes) + 8
        ws.cell(row=total_row, column=1).value = "TOTAL"
        ws.cell(row=total_row, column=1).font = Font(bold=True)
        ws.cell(row=total_row, column=3).value = float(extrato.saldo_liquido())
        ws.cell(row=total_row, column=3).font = Font(bold=True)
        ws.cell(row=total_row, column=3).number_format = '#,##0.00'
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 12
        
        self.caminho_saida.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.caminho_saida)
        
        self.logger.info(f"Arquivo XLSX criado: {self.caminho_saida}")


class ExportadorXML(ExportadorBase):
    """Export transactions to XML format."""
    
    def exportar(self, extrato: Extrato) -> None:
        """Generate XML file.
        
        Args:
            extrato: Extrato object
        """
        self.logger.info(f"Gerando XML: {self.caminho_saida}")
        
        xml = '<?xml version="1.0" encoding="UTF-8"?>\n'
        xml += '<extrato>\n'
        xml += f'  <banco>{escape(extrato.banco.value)}</banco>\n'
        xml += f'  <agencia>{escape(extrato.agencia)}</agencia>\n'
        xml += f'  <conta>{escape(extrato.conta)}</conta>\n'
        xml += f'  <data_inicio>{extrato.data_inicio}</data_inicio>\n'
        xml += f'  <data_fim>{extrato.data_fim}</data_fim>\n'
        xml += f'  <resumo>\n'
        xml += f'    <total_creditos>{extrato.total_creditos()}</total_creditos>\n'
        xml += f'    <total_debitos>{extrato.total_debitos()}</total_debitos>\n'
        xml += f'    <saldo_liquido>{extrato.saldo_liquido()}</saldo_liquido>\n'
        xml += f'  </resumo>\n'
        xml += f'  <transacoes>\n'
        
        for transacao in extrato.transacoes:
            xml += f'    <transacao>\n'
            xml += f'      <data>{transacao.data}</data>\n'
            xml += f'      <descricao>{escape(transacao.descricao)}</descricao>\n'
            xml += f'      <valor>{transacao.valor}</valor>\n'
            xml += f'      <tipo>{transacao.tipo.value}</tipo>\n'
            xml += f'    </transacao>\n'
        
        xml += f'  </transacoes>\n'
        xml += '</extrato>\n'
        
        self.caminho_saida.parent.mkdir(parents=True, exist_ok=True)
        with open(self.caminho_saida, 'w', encoding='utf-8') as arquivo:
            arquivo.write(xml)
        
        self.logger.info(f"Arquivo XML criado: {self.caminho_saida}")


class ExportadorFactory:
    """Factory to create appropriate exporter based on format."""
    
    _exportadores = {
        'ofx': ExportadorOFX,
        'csv': ExportadorCSV,
        'json': ExportadorJSON,
        'xlsx': ExportadorXLSX,
        'xml': ExportadorXML,
    }
    
    @classmethod
    def criar_exportador(cls, formato: str, caminho_saida: Path) -> ExportadorBase:
        """Create exporter for specified format.
        
        Args:
            formato: File format (ofx, csv, json, xlsx, xml)
            caminho_saida: Output file path
            
        Returns:
            Exporter instance
            
        Raises:
            ValueError: If format is not supported
        """
        formato = formato.lower()
        if formato not in cls._exportadores:
            raise ValueError(f"Formato não suportado: {formato}")
        
        exporter_class = cls._exportadores[formato]
        return exporter_class(caminho_saida)
    
    @classmethod
    def formatos_suportados(cls) -> list[str]:
        """Get list of supported formats."""
        return list(cls._exportadores.keys())
    
    @classmethod
    def registrar_exportador(cls, formato: str, exporter_class):
        """Register custom exporter for format.
        
        Args:
            formato: Format identifier
            exporter_class: Exporter class
        """
        cls._exportadores[formato.lower()] = exporter_class
        logger.info(f"Exportador registrado para formato: {formato}")
